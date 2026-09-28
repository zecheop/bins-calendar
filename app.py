from __future__ import annotations

import base64
import calendar
import hashlib
import hmac
import json
import os
import random
import re
import threading
from copy import deepcopy
from datetime import date, datetime, time as datetime_time, timedelta
from io import BytesIO
from pathlib import Path
from typing import Any
from xml.etree import ElementTree
from zoneinfo import ZoneInfo

import requests
from dotenv import load_dotenv
from flask import Flask, jsonify, request, send_from_directory
from google.api_core.exceptions import GoogleAPIError
from google.auth.transport import requests as google_requests
from google.oauth2 import id_token, service_account
from googleapiclient.discovery import build
from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent
STATIC_DIR = BASE_DIR / "static"
VIDEOS_DIR = BASE_DIR / "videos"
DATA_DIR = BASE_DIR / "data"
POSTIT_FILE = DATA_DIR / "postit.json"
CALENDAR_OVERRIDE_FILE = DATA_DIR / "calendar_overrides.json"
FROZEN_SNAPSHOT_FILE = DATA_DIR / "frozen_sheet_snapshot.json"
SONGBOOK_CUSTOM_FILE = DATA_DIR / "songbook_custom.json"
SONGBOOK_OVERRIDE_FILE = DATA_DIR / "songbook_overrides.json"

load_dotenv(BASE_DIR.parent / ".env")
load_dotenv(BASE_DIR / ".env", override=True)

TIMEZONE = ZoneInfo("Asia/Seoul")
REQUEST_TIMEOUT = 25
CACHE_TTL_SECONDS = 75
YOUTUBE_CACHE_TTL_DEFAULT_SECONDS = 4 * 60 * 60
YOUTUBE_CACHE_TTL_PREMIERE_SECONDS = 5 * 60
YOUTUBE_CACHE_TTL_EVENING_SECONDS = 15 * 60
CHZZK_API_BASE = "https://api.chzzk.naver.com"
NAVER_CAFE_API_BASE = "https://apis.naver.com/cafe-web/cafe-boardlist-api/v1"
MONTH_SHEET_RE = re.compile(r"^(?P<yy>\d{2})년\s*(?P<month>\d{1,2})월$")
YOUTUBE_HYPERLINK_RE = re.compile(r'=HYPERLINK\("([^"]+)",\s*"([^"]+)"\)', re.IGNORECASE)
IMAGE_FORMULA_RE = re.compile(r'=IMAGE\("([^"]+)"(?:\s*,.*)?\)', re.IGNORECASE)
SESSION_COOKIE_NAME = "vince_calendar_session"
MONTH_COLUMNS = [2, 5, 8, 11, 14, 17, 20]
MONTH_WEEK_ROWS = [12, 20, 28, 36, 44]
CONTENT_ROW_OFFSETS = [2, 4, 6]
LEGEND_CELLS = ["Y53", "Y54", "Y55", "Y56", "Y57", "Y58"]
LEGEND_KEYS = ["warmup", "main", "afterglow", "gap_collab", "ad", "plus"]
BRAND_ENGLISH_NAME = "Vince Calendar"
BRAND_NAME = "빈스 캘린더"
PROFILE_AVATAR_URL = "/static/assets/vince-profile.webp"
DEFAULT_LEGEND = [
    {"key": "warmup", "label": "예열", "color": "#fff9c4"},
    {"key": "main", "label": "메인컨텐츠", "color": "#cdfcf5"},
    {"key": "afterglow", "label": "후열", "color": "#fce5cd"},
    {"key": "gap_collab", "label": "틈새 합방", "color": "#ffd1da"},
    {"key": "ad", "label": "광고", "color": "#dcc8f0"},
    {"key": "plus", "label": "멤버쉽/오프라인", "color": "#a8e6cf"},
]

GOOGLE_CLIENT_ID = os.getenv("GOOGLE_CLIENT_ID", "").strip()
SESSION_SECRET = os.getenv("SESSION_SECRET", "dev-only-secret-change-me").strip()
SPREADSHEET_ID = os.getenv("SPREADSHEET_ID", "").strip()
SPREADSHEET_EXPORT_URL = (
    os.getenv("SPREADSHEET_EXPORT_URL", "").strip()
    or (f"https://docs.google.com/spreadsheets/d/{SPREADSHEET_ID}/export?format=xlsx" if SPREADSHEET_ID else "")
)
GOOGLE_SERVICE_ACCOUNT_FILE = os.getenv("GOOGLE_SERVICE_ACCOUNT_FILE", "").strip()
# 빈스는 치지직이 아니라 SOOP(숲)에서 방송함.
CHZZK_CHANNEL_ID = os.getenv("CHZZK_CHANNEL_ID", "").strip()
SOOP_BJID = os.getenv("SOOP_BJID", "psb010203").strip()
SOOP_STATION_URL = os.getenv("SOOP_STATION_URL", "https://www.sooplive.com/station/psb010203").strip()
SOOP_LIVE_API_URL = "https://live.afreecatv.com/afreeca/player_live_api.php"
YOUTUBE_CHANNEL_ID = os.getenv("YOUTUBE_CHANNEL_ID", "UCXoZBh4NsEHDDzpkqCHGGdA").strip()
YOUTUBE_CHANNEL_URL = os.getenv(
    "YOUTUBE_CHANNEL_URL",
    "https://www.youtube.com/channel/UCXoZBh4NsEHDDzpkqCHGGdA",
).strip()
CAFE_URL = os.getenv("CAFE_URL", "https://cafe.naver.com/binssss").strip()
NAVER_CAFE_ID = os.getenv("NAVER_CAFE_ID", "31114233").strip()
NAVER_CAFE_NOTICE_MENU_ID = os.getenv("NAVER_CAFE_NOTICE_MENU_ID", "").strip()
HOST = os.getenv("HOST", "0.0.0.0").strip()
PORT = int(os.getenv("PORT", "8026").strip())
ALLOWED_EDITOR_EMAILS = {
    item.strip().lower()
    for item in os.getenv("ALLOWED_EDITOR_EMAILS", "").split(",")
    if item.strip()
}
ADMIN_EMAILS = {
    item.strip().lower()
    for item in os.getenv("ADMIN_EMAILS", "wjddndj2@gmail.com").split(",")
    if item.strip()
}

snapshot_cache: dict[str, Any] = {"loaded_at": None, "data": None}
snapshot_lock = threading.Lock()
live_status_cache: dict[str, Any] = {"loaded_at": None, "data": None}
live_status_lock = threading.Lock()
youtube_cache: dict[str, Any] = {"loaded_at": None, "data": None}
youtube_cache_lock = threading.Lock()
cafe_notice_cache: dict[str, Any] = {"loaded_at": None, "data": None}
cafe_notice_lock = threading.Lock()
calendar_override_lock = threading.RLock()
songbook_custom_lock = threading.RLock()


app = Flask(__name__, static_folder=str(STATIC_DIR), static_url_path="/static")


def now_kst() -> datetime:
    return datetime.now(TIMEZONE)


def make_http_session() -> requests.Session:
    session = requests.Session()
    session.headers.update(
        {
            "User-Agent": (
                "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/146.0.0.0 Safari/537.36"
            )
        }
    )
    return session


http_session = make_http_session()


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.astimezone(TIMEZONE).strftime("%Y-%m-%d %H:%M")
    if isinstance(value, datetime_time):
        return value.strftime("%H:%M")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_hex(raw: str) -> str:
    text = str(raw or "").strip().lower()
    if not text:
        return ""
    if not text.startswith("#"):
        text = f"#{text}"
    if len(text) == 9:
        text = f"#{text[-6:]}"
    return text


def color_to_hex(color: Any) -> str:
    if color is None:
        return ""
    if getattr(color, "type", None) == "rgb" and getattr(color, "rgb", None):
        return normalize_hex(str(color.rgb))
    return ""


def parse_hyperlink_formula(value: Any) -> dict[str, str] | None:
    text = normalize_text(value)
    if not text:
        return None
    match = YOUTUBE_HYPERLINK_RE.search(text)
    if not match:
        return None
    return {"url": match.group(1), "title": match.group(2)}


def parse_image_formula(value: Any) -> str:
    text = normalize_text(value)
    if not text:
        return ""
    match = IMAGE_FORMULA_RE.search(text)
    if not match:
        return ""
    return normalize_text(match.group(1))


def extract_comment_text(cell: Any) -> str:
    comment = getattr(cell, "comment", None)
    if not comment or not getattr(comment, "text", None):
        return ""
    return normalize_text(comment.text)


def summarize_note_label(note_text: str) -> str:
    normalized = normalize_text(note_text)
    if not normalized:
        return ""

    candidates = [
        token.strip()
        for token in re.split(r"[,\n/]+", normalized)
        if token.strip()
    ]
    if len(candidates) >= 2:
        return f"참여 {len(candidates)}명"
    return "더보기"


def build_session_cookie(payload: dict[str, Any]) -> str:
    encoded = base64.urlsafe_b64encode(
        json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    ).decode("utf-8").rstrip("=")
    signature = hmac.new(SESSION_SECRET.encode("utf-8"), encoded.encode("utf-8"), hashlib.sha256).hexdigest()
    return f"{encoded}.{signature}"


def parse_session_cookie(raw_cookie: str | None) -> dict[str, Any] | None:
    if not raw_cookie or "." not in raw_cookie:
        return None
    encoded, signature = raw_cookie.rsplit(".", 1)
    expected = hmac.new(SESSION_SECRET.encode("utf-8"), encoded.encode("utf-8"), hashlib.sha256).hexdigest()
    if not hmac.compare_digest(signature, expected):
        return None
    try:
        payload_json = base64.urlsafe_b64decode(encoded + "=" * (-len(encoded) % 4)).decode("utf-8")
        payload = json.loads(payload_json)
    except Exception:
        return None
    expires_at = str(payload.get("expiresAt") or "").strip()
    if expires_at:
        try:
            if datetime.fromisoformat(expires_at) <= now_kst():
                return None
        except ValueError:
            return None
    return payload


def current_session() -> dict[str, Any] | None:
    return parse_session_cookie(request.cookies.get(SESSION_COOKIE_NAME))


def auth_state() -> dict[str, Any]:
    session = current_session()
    if not session:
        return {
            "signedIn": False,
            "canEdit": False,
            "canPostitEdit": False,
            "canAdmin": False,
            "displayName": "",
            "email": "",
        }
    email = str(session.get("email") or "").strip().lower()
    return {
        "signedIn": True,
        "canEdit": email in ALLOWED_EDITOR_EMAILS,
        "canPostitEdit": email in ALLOWED_EDITOR_EMAILS,
        "canAdmin": email in ADMIN_EMAILS,
        "displayName": str(session.get("name") or "").strip(),
        "email": email,
    }


def editor_required() -> tuple[bool, str]:
    session = current_session()
    if not session:
        return False, "로그인이 필요합니다."
    email = str(session.get("email") or "").strip().lower()
    if email not in ALLOWED_EDITOR_EMAILS:
        return False, "편집 권한이 없는 이메일입니다."
    return True, ""


def writer_enabled() -> bool:
    return bool(GOOGLE_SERVICE_ACCOUNT_FILE and Path(GOOGLE_SERVICE_ACCOUNT_FILE).exists())


def postit_editor_required() -> tuple[bool, str]:
    session = current_session()
    if not session:
        return False, "로그인이 필요합니다."
    email = str(session.get("email") or "").strip().lower()
    if email not in ALLOWED_EDITOR_EMAILS:
        return False, "편집 권한이 없는 이메일입니다."
    return True, ""


def read_only_mode_reason() -> str:
    if writer_enabled():
        return ""
    return "Google 서비스 계정이 없어도 사이트 안 편집은 가능하지만, 기존 시트로의 자동 동기화는 비활성화되어 있습니다."


def google_login_enabled() -> bool:
    return bool(GOOGLE_CLIENT_ID and SESSION_SECRET)


def get_writer_service():
    if not writer_enabled():
        return None
    credentials = service_account.Credentials.from_service_account_file(
        GOOGLE_SERVICE_ACCOUNT_FILE,
        scopes=["https://www.googleapis.com/auth/spreadsheets"],
    )
    return build("sheets", "v4", credentials=credentials, cache_discovery=False)


def fetch_workbook_bytes() -> bytes:
    response = http_session.get(SPREADSHEET_EXPORT_URL, timeout=REQUEST_TIMEOUT)
    response.raise_for_status()
    return response.content


def build_month_label(year: int, month: int) -> str:
    return f"{year}년 {month}월"


def human_month_title(year: int, month: int) -> str:
    return f"{year}년 {month:02d}월"


def format_schedule_badge(value: Any) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, (int, float)):
        numeric = float(value)
        hour = int(numeric)
        minute = 30 if abs(numeric - hour - 0.5) < 0.01 else 0
        return f"뱅온 [{hour}:{minute:02d}PM]" if minute else f"뱅온 [{hour}PM]"
    text = normalize_text(value)
    if re.fullmatch(r"\d+(?:\.5)?", text):
        return format_schedule_badge(float(text))
    return text


def derive_category_from_text(text: str, legend_lookup: dict[str, dict[str, str]]) -> str:
    normalized = text.replace(" ", "")
    for item in legend_lookup.values():
        short_label = normalize_text(item["label"]).replace(" ", "")
        if short_label and short_label in normalized:
            return item["key"]
    fallback_map = {
        "예열": "warmup",
        "메인": "main",
        "후열": "afterglow",
        "틈새": "gap_collab",
        "광고": "ad",
        "플러스": "plus",
        "멤버쉽": "plus",
        "오프라인": "plus",
    }
    for token, key in fallback_map.items():
        if token in normalized:
            return key
    return ""


def extract_legend(ws) -> list[dict[str, str]]:
    legend: list[dict[str, str]] = []
    for key, coord, default in zip(LEGEND_KEYS, LEGEND_CELLS, DEFAULT_LEGEND):
        cell = ws[coord]
        label = normalize_text(cell.value) or default["label"]
        if key == "plus":
            label = default["label"]
        color = color_to_hex(cell.fill.fgColor) or default["color"]
        legend.append({"key": key, "label": label, "color": color})
    return legend


def parse_month_sheet(ws, year: int, month: int) -> dict[str, Any]:
    legend = extract_legend(ws)
    legend_lookup = {item["key"]: item for item in legend}
    color_lookup = {normalize_hex(item["color"]): item for item in legend}
    weeks: list[list[dict[str, Any] | None]] = []
    day_items: list[dict[str, Any]] = []
    sheet_videos: list[dict[str, str]] = []
    has_entries = False

    for row_index in range(12, 60):
        link = parse_hyperlink_formula(ws[f"Y{row_index}"].value)
        if link:
            sheet_videos.append(link)

    for base_row in MONTH_WEEK_ROWS:
        week: list[dict[str, Any] | None] = []
        for weekday_index, col_start in enumerate(MONTH_COLUMNS):
            day_raw = ws.cell(base_row, col_start).value
            day_number = int(float(day_raw)) if isinstance(day_raw, (int, float)) else None
            if not day_number:
                week.append(None)
                continue

            time_cell = ws.cell(base_row, col_start + 1)
            time_label = format_schedule_badge(time_cell.value)
            is_off = normalize_text(time_label) in {"휴뱅", "휴방"}
            entries: list[dict[str, Any]] = []
            day_note_parts: list[str] = []

            day_cell_note = extract_comment_text(ws.cell(base_row, col_start))
            time_cell_note = extract_comment_text(time_cell)
            for note_text in [day_cell_note, time_cell_note]:
                if note_text and note_text not in day_note_parts:
                    day_note_parts.append(note_text)

            for offset in CONTENT_ROW_OFFSETS:
                entry_cell = ws.cell(base_row + offset, col_start)
                image_url = parse_image_formula(entry_cell.value)
                entry_text = normalize_text(entry_cell.value)
                if image_url and entry_text.startswith("="):
                    entry_text = ""
                entry_note = extract_comment_text(entry_cell)
                if not entry_text and not image_url:
                    continue
                fill_hex = normalize_hex(color_to_hex(entry_cell.fill.fgColor))
                legend_match = color_lookup.get(fill_hex)
                category_key = legend_match["key"] if legend_match else derive_category_from_text(entry_text, legend_lookup)
                category_label = legend_lookup.get(category_key, {}).get("label", "")
                category_color = legend_lookup.get(category_key, {}).get("color", "")
                entries.append(
                    {
                        "text": entry_text or "이미지 일정",
                        "categoryKey": category_key,
                        "categoryLabel": category_label,
                        "categoryColor": category_color,
                        "noteText": entry_note,
                        "noteLabel": summarize_note_label(entry_note),
                        "imageUrl": image_url,
                    }
                )

            has_entries = has_entries or bool(time_label or entries)
            day_note_text = "\n\n".join(day_note_parts)
            day_payload = {
                "day": day_number,
                "timeLabel": time_label,
                "timeSource": "sheet" if time_label else "",
                "isOff": is_off,
                "entries": entries,
                "weekdayIndex": weekday_index,
                "baseRow": base_row,
                "baseColumn": col_start,
                "noteText": day_note_text,
                "noteLabel": summarize_note_label(day_note_text),
                "showStatusImage": bool(is_off),
            }
            week.append(day_payload)
            day_items.append(day_payload)
        weeks.append(week)

    return {
        "key": f"{year}-{month:02d}",
        "sheetName": ws.title,
        "title": human_month_title(year, month),
        "year": year,
        "month": month,
        "weeks": weeks,
        "days": day_items,
        "legend": legend,
        "sheetVideos": sheet_videos[:6],
        "hasEntries": has_entries,
        "source": "sheet",
    }


def parse_songbook_sheet(ws) -> dict[str, Any]:
    public_link = normalize_text(ws["H5"].value)
    items: list[dict[str, Any]] = []
    for row in range(6, min(ws.max_row, 300) + 1):
        number = ws[f"B{row}"].value
        title = normalize_text(ws[f"C{row}"].value)
        if number in (None, "") or not title:
            continue
        original = normalize_text(ws[f"D{row}"].value)
        if original == '"':
            original = ""
        artist = normalize_text(ws[f"E{row}"].value)
        category = normalize_text(ws[f"F{row}"].value)
        items.append(
            {
                "id": f"sheet-{row}",
                "number": int(float(number)) if isinstance(number, (int, float)) else normalize_text(number),
                "title": title,
                "originalTitle": original,
                "artist": artist,
                "category": category,
                "source": "sheet",
            }
        )
    return {
        "notice": normalize_text(ws["B2"].value),
        "publicLink": public_link,
        "items": items,
        "categories": sorted({item["category"] for item in items if item["category"]}),
        "artists": sorted({item["artist"] for item in items if item["artist"]}),
    }


def parse_song_number(value: Any) -> int | None:
    text = normalize_text(value)
    if not text:
        return None
    if re.fullmatch(r"\d+", text):
        try:
            return int(text)
        except ValueError:
            return None
    return None


def sanitize_songbook_item(raw: Any, fallback_number: str = "", fallback_id: str = "") -> dict[str, Any] | None:
    item_id = normalize_text((raw or {}).get("id"))[:80] or fallback_id
    title = normalize_text((raw or {}).get("title"))[:120]
    original_title = normalize_text((raw or {}).get("originalTitle"))[:120]
    if original_title == '"':
        original_title = ""
    artist = normalize_text((raw or {}).get("artist"))[:80]
    category = normalize_text((raw or {}).get("category"))[:40]
    number = normalize_text((raw or {}).get("number"))[:20] or fallback_number
    if not title:
        return None
    return {
        "id": item_id or "",
        "number": number or "",
        "title": title,
        "originalTitle": original_title,
        "artist": artist,
        "category": category,
        "source": normalize_text((raw or {}).get("source")) or "custom",
    }


def sanitize_songbook_items(items: Any) -> list[dict[str, Any]]:
    sanitized: list[dict[str, Any]] = []
    raw_items = items if isinstance(items, list) else []
    for index, raw in enumerate(raw_items[:1000], start=1):
        fallback_id = normalize_text((raw or {}).get("id")) or f"song-{index}"
        item = sanitize_songbook_item(raw, fallback_id=fallback_id)
        if item:
            sanitized.append(item)
    return sanitized


def load_songbook_custom_items() -> list[dict[str, Any]]:
    with songbook_custom_lock:
        if not SONGBOOK_CUSTOM_FILE.exists():
            return []
        try:
            payload = json.loads(SONGBOOK_CUSTOM_FILE.read_text(encoding="utf-8"))
        except Exception:
            return []
        return sanitize_songbook_items((payload.get("items") if isinstance(payload, dict) else []) or [])


def load_songbook_override_items() -> list[dict[str, Any]] | None:
    with songbook_custom_lock:
        if not SONGBOOK_OVERRIDE_FILE.exists():
            return None
        try:
            payload = json.loads(SONGBOOK_OVERRIDE_FILE.read_text(encoding="utf-8"))
        except Exception:
            return None
        return sanitize_songbook_items((payload.get("items") if isinstance(payload, dict) else []) or [])


def save_songbook_override_items(items: Any) -> list[dict[str, Any]]:
    with songbook_custom_lock:
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        sanitized = sanitize_songbook_items(items)
        SONGBOOK_OVERRIDE_FILE.write_text(
            json.dumps({"items": sanitized, "updatedAt": now_kst().isoformat()}, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        return sanitized


def build_songbook_payload(snapshot_songbook: dict[str, Any]) -> dict[str, Any]:
    override_items = load_songbook_override_items()
    if override_items is not None:
        items = override_items
    else:
        base_items = list(snapshot_songbook.get("items") or [])
        custom_items = load_songbook_custom_items()
        items = base_items + custom_items
    return {
        "notice": snapshot_songbook.get("notice", ""),
        "publicLink": snapshot_songbook.get("publicLink", ""),
        "items": items,
        "categories": sorted({normalize_text(item.get("category")) for item in items if normalize_text(item.get("category"))}),
        "artists": sorted({normalize_text(item.get("artist")) for item in items if normalize_text(item.get("artist"))}),
    }


def choose_initial_month(months: list[dict[str, Any]]) -> dict[str, int]:
    current = now_kst()
    return {"year": current.year, "month": current.month}


def sanitize_override_entries(entries: Any) -> list[dict[str, str]]:
    sanitized: list[dict[str, str]] = []
    for raw in (entries if isinstance(entries, list) else [])[:5]:
        text = normalize_text((raw or {}).get("text"))[:120]
        category_key = normalize_text((raw or {}).get("categoryKey"))
        note_text = normalize_text((raw or {}).get("noteText"))[:120]
        if not text and not note_text:
            continue
        sanitized.append(
            {
                "text": text or "기타 메모",
                "categoryKey": category_key,
                "noteText": note_text,
                "noteLabel": summarize_note_label(note_text) if note_text else "",
            }
        )
    return sanitized


def sanitize_day_override(raw: Any) -> dict[str, Any]:
    time_label = normalize_text(format_schedule_badge((raw or {}).get("timeLabel")))
    if time_label in {"휴방", "휴뱅"}:
        time_label = "휴뱅"
    time_source = normalize_text((raw or {}).get("timeSource")) or "manual"
    day_note_text = normalize_text((raw or {}).get("noteText"))[:1200]
    if time_source not in {"manual", "auto-live"}:
        time_source = "manual"
    return {
        "timeLabel": time_label,
        "timeSource": time_source,
        "entries": sanitize_override_entries((raw or {}).get("entries")),
        "noteText": day_note_text,
        "noteLabel": summarize_note_label(day_note_text) if day_note_text else "",
        "updatedAt": normalize_text((raw or {}).get("updatedAt")) or now_kst().isoformat(),
    }


def load_calendar_overrides() -> dict[str, Any]:
    with calendar_override_lock:
        if not CALENDAR_OVERRIDE_FILE.exists():
            return {"months": {}}
        try:
            payload = json.loads(CALENDAR_OVERRIDE_FILE.read_text(encoding="utf-8"))
        except Exception:
            return {"months": {}}

        months: dict[str, Any] = {}
        for raw_key, raw_month in (payload.get("months") or {}).items():
            month_key = normalize_text(raw_key)
            if not re.fullmatch(r"\d{4}-\d{2}", month_key):
                continue
            days: dict[str, Any] = {}
            for raw_day, day_payload in ((raw_month or {}).get("days") or {}).items():
                try:
                    day_number = int(raw_day)
                except (TypeError, ValueError):
                    continue
                if 1 <= day_number <= 31:
                    days[str(day_number)] = sanitize_day_override(day_payload)
            if days:
                months[month_key] = {"days": days}
        return {"months": months}


def save_calendar_overrides(payload: Any) -> dict[str, Any]:
    with calendar_override_lock:
        sanitized = {"months": {}}
        for raw_key, raw_month in ((payload or {}).get("months") or {}).items():
            month_key = normalize_text(raw_key)
            if not re.fullmatch(r"\d{4}-\d{2}", month_key):
                continue
            days: dict[str, Any] = {}
            for raw_day, day_payload in ((raw_month or {}).get("days") or {}).items():
                try:
                    day_number = int(raw_day)
                except (TypeError, ValueError):
                    continue
                if 1 <= day_number <= 31:
                    days[str(day_number)] = sanitize_day_override(day_payload)
            if days:
                sanitized["months"][month_key] = {"days": days}

        DATA_DIR.mkdir(parents=True, exist_ok=True)
        CALENDAR_OVERRIDE_FILE.write_text(
            json.dumps(sanitized, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        return sanitized


def update_calendar_day_override(
    year: int,
    month: int,
    day: int,
    time_text: str,
    entries: list[dict[str, str]],
    day_note_text: str = "",
    *,
    time_source: str = "manual",
) -> dict[str, Any]:
    payload = load_calendar_overrides()
    month_key = f"{year}-{month:02d}"
    month_payload = payload["months"].setdefault(month_key, {"days": {}})
    month_payload["days"][str(day)] = sanitize_day_override(
        {
            "timeLabel": time_text,
            "timeSource": time_source,
            "entries": entries,
            "noteText": day_note_text,
            "updatedAt": now_kst().isoformat(),
        }
    )
    return save_calendar_overrides(payload)


def build_blank_month(year: int, month: int, legend: list[dict[str, str]] | None = None) -> dict[str, Any]:
    current_legend = deepcopy(legend or DEFAULT_LEGEND)
    first_weekday = (date(year, month, 1).weekday() + 1) % 7
    days_in_month = calendar.monthrange(year, month)[1]
    weeks: list[list[dict[str, Any] | None]] = []
    current_week: list[dict[str, Any] | None] = [None] * 7
    day_items: list[dict[str, Any]] = []
    weekday_index = first_weekday

    for day_number in range(1, days_in_month + 1):
        day_payload = {
            "day": day_number,
            "timeLabel": "",
            "timeSource": "",
            "isOff": False,
            "entries": [],
            "weekdayIndex": weekday_index,
            "baseRow": 0,
            "baseColumn": 0,
            "noteText": "",
            "noteLabel": "",
            "showStatusImage": False,
        }
        current_week[weekday_index] = day_payload
        day_items.append(day_payload)

        if weekday_index == 6 or day_number == days_in_month:
            weeks.append(current_week)
            current_week = [None] * 7
        weekday_index = (weekday_index + 1) % 7

    return {
        "key": f"{year}-{month:02d}",
        "sheetName": "",
        "title": human_month_title(year, month),
        "year": year,
        "month": month,
        "weeks": weeks,
        "days": day_items,
        "legend": current_legend,
        "sheetVideos": [],
        "hasEntries": False,
        "source": "local",
    }


def apply_calendar_overrides_to_month(month_data: dict[str, Any], overrides: dict[str, Any]) -> dict[str, Any]:
    legend_lookup = {item["key"]: item for item in month_data.get("legend") or DEFAULT_LEGEND}
    month_override = ((overrides.get("months") or {}).get(month_data["key"]) or {}).get("days") or {}
    day_lookup = {int(item["day"]): item for item in month_data.get("days") or []}

    for raw_day, override in month_override.items():
        try:
            day_number = int(raw_day)
        except (TypeError, ValueError):
            continue
        day_payload = day_lookup.get(day_number)
        if not day_payload:
            continue

        time_label = normalize_text((override or {}).get("timeLabel"))
        if time_label in {"휴방", "휴뱅"}:
            time_label = "휴뱅"
        rendered_entries: list[dict[str, Any]] = []
        for entry in (override or {}).get("entries") or []:
            entry_text = normalize_text((entry or {}).get("text"))
            if not entry_text:
                continue
            category_key = normalize_text((entry or {}).get("categoryKey"))
            note_text = normalize_text((entry or {}).get("noteText"))
            legend_item = legend_lookup.get(category_key, {})
            rendered_entries.append(
                {
                    "text": entry_text,
                    "categoryKey": category_key,
                    "categoryLabel": legend_item.get("label", ""),
                    "categoryColor": legend_item.get("color", ""),
                    "noteText": note_text,
                    "noteLabel": normalize_text((entry or {}).get("noteLabel")) or (summarize_note_label(note_text) if note_text else ""),
                    "imageUrl": "",
                }
            )

        day_payload["timeLabel"] = time_label
        day_payload["timeSource"] = normalize_text((override or {}).get("timeSource")) or "manual"
        day_payload["isOff"] = time_label == "휴뱅"
        day_payload["entries"] = rendered_entries
        override_note_text = normalize_text((override or {}).get("noteText"))
        day_payload["noteText"] = override_note_text
        day_payload["noteLabel"] = normalize_text((override or {}).get("noteLabel")) or (summarize_note_label(override_note_text) if override_note_text else "")
        day_payload["localOverride"] = True
    return month_data


def finalize_month_payload(month_data: dict[str, Any]) -> dict[str, Any]:
    has_entries = False
    for day_payload in month_data.get("days") or []:
        time_label = normalize_text(day_payload.get("timeLabel"))
        day_payload["isOff"] = time_label in {"휴방", "휴뱅"}
        day_payload["showStatusImage"] = bool(day_payload["isOff"])
        has_entries = has_entries or bool(time_label or day_payload.get("entries"))
    month_data["hasEntries"] = has_entries
    return month_data


def build_calendar_bounds(months_meta: list[dict[str, Any]], overrides: dict[str, Any]) -> dict[str, int]:
    current_year = now_kst().year
    years = {current_year - 1, current_year, current_year + 1, current_year + 2}
    years.update(item["year"] for item in months_meta if item.get("year"))
    for month_key in (overrides.get("months") or {}):
        try:
            years.add(int(month_key.split("-", 1)[0]))
        except (IndexError, ValueError):
            continue
    return {"minYear": min(years), "maxYear": max(years)}


def build_bootstrap_months_meta(snapshot: dict[str, Any], overrides: dict[str, Any]) -> list[dict[str, Any]]:
    month_map = {item["key"]: dict(item) for item in snapshot["monthsMeta"]}
    for month_key in (overrides.get("months") or {}):
        if month_key in month_map:
            month_map[month_key]["hasEntries"] = True
            continue
        year_text, month_text = month_key.split("-", 1)
        year = int(year_text)
        month = int(month_text)
        month_map[month_key] = {
            "key": month_key,
            "sheetName": "",
            "title": human_month_title(year, month),
            "year": year,
            "month": month,
            "hasEntries": True,
        }
    return sorted(month_map.values(), key=lambda item: (item["year"], item["month"]))


def get_month_payload(year: int, month: int) -> dict[str, Any]:
    if not (1 <= month <= 12):
        raise ValueError("month must be between 1 and 12")

    snapshot = load_snapshot()
    overrides = load_calendar_overrides()
    month_key = f"{year}-{month:02d}"
    base_month = snapshot["monthsByKey"].get(month_key)

    if base_month:
        month_data = build_blank_month(year, month, base_month.get("legend") or snapshot.get("legend") or DEFAULT_LEGEND)
        month_data["sheetName"] = base_month.get("sheetName", "")
        month_data["sheetVideos"] = deepcopy(base_month.get("sheetVideos") or [])
        month_data["source"] = "sheet"
        day_lookup = {int(item["day"]): item for item in month_data.get("days") or []}
        for source_day in base_month.get("days") or []:
            target_day = day_lookup.get(int(source_day["day"]))
            if target_day:
                target_day.update(deepcopy(source_day))
    else:
        month_data = build_blank_month(year, month, snapshot.get("legend") or DEFAULT_LEGEND)

    apply_calendar_overrides_to_month(month_data, overrides)
    return finalize_month_payload(month_data)


def read_frozen_snapshot() -> dict[str, Any] | None:
    if not FROZEN_SNAPSHOT_FILE.exists():
        return None
    try:
        snapshot = json.loads(FROZEN_SNAPSHOT_FILE.read_text(encoding="utf-8"))
    except Exception:
        return None
    required_keys = {"generatedAt", "monthsMeta", "monthsByKey", "legend", "songbook", "initial"}
    if not required_keys.issubset(snapshot):
        return None
    return snapshot


def write_frozen_snapshot(snapshot: dict[str, Any]) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    FROZEN_SNAPSHOT_FILE.write_text(
        json.dumps(snapshot, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def build_snapshot_from_workbook() -> dict[str, Any]:
    workbook = load_workbook(BytesIO(fetch_workbook_bytes()), data_only=False)
    months_meta: list[dict[str, Any]] = []
    months_by_key: dict[str, dict[str, Any]] = {}
    songbook = {"notice": "", "publicLink": "", "items": []}

    for sheet_name in workbook.sheetnames:
        match = MONTH_SHEET_RE.match(sheet_name)
        if match:
            year = 2000 + int(match.group("yy"))
            month = int(match.group("month"))
            parsed = parse_month_sheet(workbook[sheet_name], year, month)
            months_by_key[parsed["key"]] = parsed
            months_meta.append(
                {
                    "key": parsed["key"],
                    "sheetName": parsed["sheetName"],
                    "title": parsed["title"],
                    "year": parsed["year"],
                    "month": parsed["month"],
                    "hasEntries": parsed["hasEntries"],
                }
            )
            continue
        # TODO: 빈스 전용 스프레드시트가 생기면 실제 노래책 시트 이름으로 교체
        if sheet_name == "노래책":
            songbook = parse_songbook_sheet(workbook[sheet_name])

    months_meta.sort(key=lambda item: (item["year"], item["month"]))
    initial = choose_initial_month(months_meta.copy())
    initial_key = f"{initial['year']}-{initial['month']:02d}"
    initial_month = months_by_key.get(initial_key)
    legend = initial_month["legend"] if initial_month else DEFAULT_LEGEND

    return {
        "generatedAt": now_kst().isoformat(),
        "monthsMeta": months_meta,
        "monthsByKey": months_by_key,
        "legend": legend,
        "songbook": songbook,
        "initial": initial,
    }


def load_snapshot(force: bool = False) -> dict[str, Any]:
    with snapshot_lock:
        frozen_snapshot = read_frozen_snapshot()
        if frozen_snapshot:
            snapshot_cache["loaded_at"] = now_kst()
            snapshot_cache["data"] = frozen_snapshot
            return frozen_snapshot

        loaded_at = snapshot_cache["loaded_at"]
        if not force and loaded_at and snapshot_cache["data"]:
            if now_kst() - loaded_at < timedelta(seconds=CACHE_TTL_SECONDS):
                return snapshot_cache["data"]

        snapshot = build_snapshot_from_workbook()
        write_frozen_snapshot(snapshot)
        snapshot_cache["loaded_at"] = now_kst()
        snapshot_cache["data"] = snapshot
        return snapshot


def format_live_start_badge(start_at: datetime) -> str:
    local_time = start_at.astimezone(TIMEZONE)
    meridiem = "AM" if local_time.hour < 12 else "PM"
    hour = local_time.hour % 12 or 12
    if local_time.minute:
        return f"뱅온 [{hour}:{local_time.minute:02d}{meridiem}]"
    return f"뱅온 [{hour}{meridiem}]"


def maybe_capture_live_start(live_payload: dict[str, Any]) -> None:
    if not live_payload.get("isLive") or not live_payload.get("startAt"):
        return
    # 스프레드시트가 아직 없는 채널(예: 빈스)은 시트 기반 캘린더 동기화 자체가
    # 불가능하므로, 여기서 조용히 건너뛴다(에디터가 수동으로 일정을 등록한다).
    if not SPREADSHEET_EXPORT_URL:
        return

    try:
        start_at = datetime.fromisoformat(str(live_payload["startAt"])).astimezone(TIMEZONE)
    except ValueError:
        return

    month_key = f"{start_at.year}-{start_at.month:02d}"
    snapshot = load_snapshot()
    if month_key in snapshot["monthsByKey"]:
        return

    month_data = get_month_payload(start_at.year, start_at.month)
    target_day = next((item for item in month_data["days"] if int(item["day"]) == start_at.day), None)
    if not target_day:
        return

    existing_time = normalize_text(target_day.get("timeLabel"))
    if existing_time:
        return

    update_calendar_day_override(
        start_at.year,
        start_at.month,
        start_at.day,
        format_live_start_badge(start_at),
        [
            {
                "text": normalize_text(entry.get("text")),
                "categoryKey": normalize_text(entry.get("categoryKey")),
            }
            for entry in (target_day.get("entries") or [])
            if normalize_text(entry.get("text"))
        ],
        time_source="auto-live",
    )


def fetch_soop_live_status() -> dict[str, Any]:
    # SOOP(구 아프리카TV)의 비공식 플레이어 API. 공식 문서는 없고, 이 엔드포인트가
    # 방송 중 여부·제목·채팅 서버 정보를 담아 돌려준다는 것은 커뮤니티 역공학으로
    # 확인된 내용이다(static/game/soop-chat.js 참고). BTIME은 절대시각이 아니라
    # "방송 시작 후 경과 초"라서, 시작 시각은 지금 시각에서 그만큼 빼서 계산한다.
    response = http_session.post(
        f"{SOOP_LIVE_API_URL}?bjid={SOOP_BJID}",
        data={
            "bid": SOOP_BJID,
            "bno": "",
            "type": "live",
            "confirm_adult": "false",
            "player_type": "html5",
            "mode": "landing",
            "from_api": "0",
            "pwd": "",
            "stream_type": "common",
            "quality": "HD",
        },
        timeout=REQUEST_TIMEOUT,
    )
    channel_payload = response.json().get("CHANNEL", {})
    live_now = int(channel_payload.get("RESULT") or 0) == 1

    start_at = ""
    if live_now:
        try:
            elapsed_seconds = int(channel_payload.get("BTIME") or 0)
            start_at = (now_kst() - timedelta(seconds=elapsed_seconds)).isoformat()
        except (TypeError, ValueError):
            start_at = ""

    payload = {
        "channelId": SOOP_BJID,
        "channelName": normalize_text(channel_payload.get("BJNICK")) or "빈스",
        "channelUrl": SOOP_STATION_URL,
        "avatarUrl": "",
        "bio": "알다가도 모를 까마귀",
        "followerCount": 0,
        "isLive": live_now,
        "liveTitle": normalize_text(channel_payload.get("TITLE")),
        "startAt": start_at,
        "closeAt": "",
        "lastStartAt": start_at,
        "lastCloseAt": "",
        "available": True,
        "stale": False,
        "message": "",
    }
    maybe_capture_live_start(payload)
    with live_status_lock:
        live_status_cache["loaded_at"] = now_kst()
        live_status_cache["data"] = payload
    return payload


def fallback_live_status(message: str, stale: bool = False) -> dict[str, Any]:
    return {
        "channelId": SOOP_BJID,
        "channelName": "빈스",
        "channelUrl": SOOP_STATION_URL,
        "avatarUrl": "",
        "bio": "알다가도 모를 까마귀",
        "followerCount": 0,
        "isLive": False,
        "liveTitle": "",
        "startAt": "",
        "closeAt": "",
        "lastStartAt": "",
        "lastCloseAt": "",
        "available": False,
        "stale": stale,
        "message": message,
    }


def get_live_status_payload() -> dict[str, Any]:
    try:
        return fetch_soop_live_status()
    except Exception:
        with live_status_lock:
            cached = live_status_cache.get("data")
        if cached:
            stale_payload = dict(cached)
            stale_payload["stale"] = True
            stale_payload["message"] = "방송 상태를 잠시 다시 확인하는 중입니다."
            return stale_payload
        return fallback_live_status("방송 상태를 잠시 확인하지 못했습니다.")


def normalize_cafe_article_items(payload: dict[str, Any]) -> list[dict[str, Any]]:
    result = payload.get("result")
    if isinstance(result, dict):
        article_list = result.get("articleList") or []
    elif isinstance(result, list):
        article_list = result
    else:
        article_list = []
    normalized: list[dict[str, Any]] = []
    for entry in article_list:
        item = (entry or {}).get("item") or {}
        if item:
            normalized.append(item)
    return normalized


def build_cafe_article_url(cafe_id: str, article_id: int) -> str:
    return f"https://cafe.naver.com/f-e/cafes/{cafe_id}/articles/{article_id}"


def pick_top_visible_cafe_notice(
    article_items: list[dict[str, Any]],
    pinned_article_items: list[dict[str, Any]],
    cafe_id: str,
) -> dict[str, Any] | None:
    pinned_ids = {
        int(item.get("articleId") or 0)
        for item in pinned_article_items
        if int(item.get("articleId") or 0) > 0
    }
    for item in article_items:
        article_id = int(item.get("articleId") or 0)
        title = normalize_text(item.get("subject"))
        if article_id <= 0 or not title or article_id in pinned_ids:
            continue
        if item.get("blindArticle") or item.get("delParent") or item.get("openArticle") is False:
            continue
        return {
            "articleId": article_id,
            "title": title,
            "url": build_cafe_article_url(cafe_id, article_id),
        }
    return None


def fetch_cafe_notice_item() -> dict[str, Any]:
    current = now_kst()
    with cafe_notice_lock:
        loaded_at = cafe_notice_cache.get("loaded_at")
        cached = cafe_notice_cache.get("data")
        if loaded_at and cached and current - loaded_at < timedelta(minutes=10):
            return cached

    headers = {
        "User-Agent": http_session.headers.get("User-Agent", "Mozilla/5.0"),
        "Referer": CAFE_URL,
        "Accept": "application/json, text/plain, */*",
    }
    article_list_url = f"{NAVER_CAFE_API_BASE}/cafes/{NAVER_CAFE_ID}/menus/{NAVER_CAFE_NOTICE_MENU_ID}/articles"
    pinned_list_url = f"{NAVER_CAFE_API_BASE}/cafes/{NAVER_CAFE_ID}/uparticles/menus/{NAVER_CAFE_NOTICE_MENU_ID}"

    article_response = requests.get(article_list_url, headers=headers, timeout=REQUEST_TIMEOUT)
    pinned_response = requests.get(pinned_list_url, headers=headers, timeout=REQUEST_TIMEOUT)
    article_response.raise_for_status()
    pinned_response.raise_for_status()

    article_items = normalize_cafe_article_items(article_response.json())
    pinned_items = normalize_cafe_article_items(pinned_response.json())
    item = pick_top_visible_cafe_notice(article_items, pinned_items, NAVER_CAFE_ID)
    if not item:
        raise RuntimeError("공지사항을 찾지 못했습니다.")

    with cafe_notice_lock:
        cafe_notice_cache["loaded_at"] = current
        cafe_notice_cache["data"] = item
    return item


def fetch_youtube_latest() -> list[dict[str, Any]]:
    feed_url = f"https://www.youtube.com/feeds/videos.xml?channel_id={YOUTUBE_CHANNEL_ID}"
    response = http_session.get(feed_url, timeout=REQUEST_TIMEOUT)
    response.raise_for_status()
    root = ElementTree.fromstring(response.text)
    ns = {
        "atom": "http://www.w3.org/2005/Atom",
        "yt": "http://www.youtube.com/xml/schemas/2015",
        "media": "http://search.yahoo.com/mrss/",
    }
    videos: list[dict[str, Any]] = []
    for entry in root.findall("atom:entry", ns)[:3]:
        link = entry.find("atom:link[@rel='alternate']", ns)
        url = normalize_text(link.attrib.get("href")) if link is not None else ""
        title = entry.findtext("atom:title", default="", namespaces=ns)
        published = entry.findtext("atom:published", default="", namespaces=ns)
        video_id = entry.findtext("yt:videoId", default="", namespaces=ns)
        thumbnail = entry.find("media:group/media:thumbnail", ns)
        videos.append(
            {
                "videoId": video_id,
                "title": title,
                "url": url or f"https://www.youtube.com/watch?v={video_id}",
                "publishedAt": published,
                "thumbnailUrl": thumbnail.attrib.get("url", "") if thumbnail is not None else "",
                "isShort": "/shorts/" in url or "#shorts" in title.lower(),
            }
        )
    return videos


def youtube_refresh_interval_seconds(current: datetime | None = None) -> int:
    now = current or now_kst()
    minutes = now.hour * 60 + now.minute
    if 17 * 60 + 30 <= minutes < 18 * 60 + 30:
        return YOUTUBE_CACHE_TTL_PREMIERE_SECONDS
    if 18 * 60 + 30 <= minutes < 21 * 60:
        return YOUTUBE_CACHE_TTL_EVENING_SECONDS
    return YOUTUBE_CACHE_TTL_DEFAULT_SECONDS


def build_youtube_fallback_items(snapshot: dict[str, Any]) -> list[dict[str, Any]]:
    items = normalize_youtube_items(snapshot.get("youtubeFallback") or [])
    if items:
        return items

    fallback = []
    initial_key = f"{snapshot['initial']['year']}-{snapshot['initial']['month']:02d}"
    month = snapshot["monthsByKey"].get(initial_key)
    if month:
        fallback = month["sheetVideos"][:3]
    return [
        {
            "videoId": "",
            "title": item["title"],
            "url": item["url"],
            "publishedAt": "",
            "thumbnailUrl": "",
            "isShort": "short" in item["url"],
        }
        for item in fallback
    ]


def normalize_youtube_items(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    normalized: list[dict[str, Any]] = []
    for item in items[:3]:
        title = normalize_text(item.get("title"))
        url = normalize_text(item.get("url"))
        if not title or not url:
            continue
        normalized.append(
            {
                "videoId": normalize_text(item.get("videoId")),
                "title": title,
                "url": url,
                "publishedAt": normalize_text(item.get("publishedAt")),
                "thumbnailUrl": normalize_text(item.get("thumbnailUrl")),
                "isShort": bool(item.get("isShort")),
            }
        )
    return normalized


def get_youtube_latest_payload() -> dict[str, Any]:
    current = now_kst()
    ttl_seconds = youtube_refresh_interval_seconds(current)
    with youtube_cache_lock:
        loaded_at = youtube_cache.get("loaded_at")
        cached = youtube_cache.get("data")
        if loaded_at and cached and current - loaded_at < timedelta(seconds=ttl_seconds):
            return {
                "items": normalize_youtube_items(cached),
                "stale": False,
                "cached": True,
                "nextPollMs": ttl_seconds * 1000,
            }

    try:
        items = normalize_youtube_items(fetch_youtube_latest())
        with youtube_cache_lock:
            youtube_cache["loaded_at"] = current
            youtube_cache["data"] = items
        return {
            "items": items,
            "stale": False,
            "cached": False,
            "nextPollMs": ttl_seconds * 1000,
        }
    except Exception as exc:
        with youtube_cache_lock:
            cached = youtube_cache.get("data")
        if cached:
            return {
                "items": normalize_youtube_items(cached),
                "stale": True,
                "cached": True,
                "nextPollMs": ttl_seconds * 1000,
                "warning": f"유튜브 RSS를 다시 확인하는 중이라 마지막 성공 데이터를 유지했습니다: {exc}",
            }

        snapshot = load_snapshot()
        return {
            "items": build_youtube_fallback_items(snapshot),
            "stale": True,
            "cached": False,
            "nextPollMs": ttl_seconds * 1000,
            "warning": f"유튜브 RSS 읽기에 실패해 저장된 대체 데이터를 사용했습니다: {exc}",
        }


def build_bootstrap_payload() -> dict[str, Any]:
    snapshot = load_snapshot()
    songbook = build_songbook_payload(snapshot["songbook"])
    overrides = load_calendar_overrides()
    months_meta = build_bootstrap_months_meta(snapshot, overrides)
    auth = auth_state()
    return {
        "generatedAt": snapshot["generatedAt"],
        "brandName": BRAND_NAME,
        "brandEnglishName": BRAND_ENGLISH_NAME,
        "profileAvatarUrl": PROFILE_AVATAR_URL,
        "initial": choose_initial_month(months_meta),
        "months": months_meta,
        "calendarBounds": build_calendar_bounds(months_meta, overrides),
        "legend": snapshot["legend"],
        "songbookSummary": {
            "count": len(songbook["items"]),
            "publicLink": songbook["publicLink"],
            "notice": songbook["notice"],
            "categories": songbook.get("categories", []),
            "artists": songbook.get("artists", []),
        },
        "auth": {
            "googleClientId": GOOGLE_CLIENT_ID,
            "googleEnabled": google_login_enabled(),
            "writeEnabled": writer_enabled(),
            "signedIn": auth["signedIn"],
            "canEdit": auth["canEdit"],
            "canPostitEdit": auth["canPostitEdit"],
            "canAdmin": auth["canAdmin"],
            "displayName": auth["displayName"],
            "email": auth["email"],
            "readOnlyReason": read_only_mode_reason(),
        },
        "links": {
            "youtubeChannelUrl": YOUTUBE_CHANNEL_URL,
            "chzzkChannelUrl": SOOP_STATION_URL,
            "cafeUrl": CAFE_URL,
            "spreadsheetUrl": (
                f"https://docs.google.com/spreadsheets/d/{SPREADSHEET_ID}/edit" if SPREADSHEET_ID else ""
            ),
        },
    }


def sanitize_postit_items(items: Any) -> list[dict[str, Any]]:
    sanitized: list[dict[str, Any]] = []
    raw_items = items if isinstance(items, list) else []
    for index, raw in enumerate(raw_items[:10], start=1):
        text = normalize_text((raw or {}).get("text"))
        checked = bool((raw or {}).get("checked"))
        item_id = normalize_text((raw or {}).get("id")) or f"postit-{index}"
        if not text and not checked:
            continue
        sanitized.append(
            {
                "id": item_id,
                "text": text[:120],
                "checked": checked,
            }
        )
    return sanitized


def load_postit_payload() -> dict[str, Any]:
    if not POSTIT_FILE.exists():
        return {"items": [], "updatedAt": ""}
    try:
        payload = json.loads(POSTIT_FILE.read_text(encoding="utf-8"))
    except Exception:
        return {"items": [], "updatedAt": ""}
    return {
        "items": sanitize_postit_items(payload.get("items") or []),
        "updatedAt": normalize_text(payload.get("updatedAt")),
    }


def save_postit_payload(items: Any) -> dict[str, Any]:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    payload = {
        "items": sanitize_postit_items(items),
        "updatedAt": now_kst().isoformat(),
    }
    POSTIT_FILE.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return payload


def get_month_or_404(year: int, month: int) -> dict[str, Any]:
    return get_month_payload(year, month)


def build_json_response(payload: Any, status: int = 200):
    response = jsonify(payload)
    response.status_code = status
    response.headers["Access-Control-Allow-Origin"] = "*"
    return response


def get_sheet_title_to_gid() -> dict[str, int]:
    service = get_writer_service()
    if service is None:
        return {}
    payload = service.spreadsheets().get(
        spreadsheetId=SPREADSHEET_ID,
        fields="sheets(properties(sheetId,title))",
    ).execute()
    result: dict[str, int] = {}
    for sheet in payload.get("sheets", []):
        props = sheet.get("properties", {})
        title = normalize_text(props.get("title"))
        gid = int(props.get("sheetId") or 0)
        if title:
            result[title] = gid
    return result


def hex_to_google_color(value: str) -> dict[str, float]:
    text = normalize_hex(value).lstrip("#")
    if len(text) != 6:
        text = "ffffff"
    return {
        "red": int(text[0:2], 16) / 255,
        "green": int(text[2:4], 16) / 255,
        "blue": int(text[4:6], 16) / 255,
    }


def write_cell_request(sheet_id: int, row: int, col: int, value: Any, fill_hex: str) -> dict[str, Any]:
    user_value: dict[str, Any]
    if value in (None, ""):
        user_value = {"stringValue": ""}
    elif isinstance(value, (int, float)):
        user_value = {"numberValue": float(value)}
    else:
        text = normalize_text(value)
        if re.fullmatch(r"\d+(?:\.5)?", text):
            user_value = {"numberValue": float(text)}
        else:
            user_value = {"stringValue": text}

    return {
        "repeatCell": {
            "range": {
                "sheetId": sheet_id,
                "startRowIndex": row - 1,
                "endRowIndex": row,
                "startColumnIndex": col - 1,
                "endColumnIndex": col,
            },
            "cell": {
                "userEnteredValue": user_value,
                "userEnteredFormat": {
                    "backgroundColor": hex_to_google_color(fill_hex),
                },
            },
            "fields": "userEnteredValue,userEnteredFormat.backgroundColor",
        }
    }


def update_day_on_sheet(
    sheet_name: str,
    base_row: int,
    base_col: int,
    time_text: str,
    entries: list[dict[str, str]],
):
    service = get_writer_service()
    if service is None:
        raise RuntimeError("쓰기 서비스가 설정되지 않았습니다.")

    gid_map = get_sheet_title_to_gid()
    sheet_id = gid_map.get(sheet_name)
    if sheet_id is None:
        raise RuntimeError("대상 시트의 sheetId 를 찾지 못했습니다.")

    snapshot = load_snapshot()
    legend_lookup = {item["key"]: item for item in snapshot["legend"]}

    requests_payload = []
    normalized_time = normalize_text(time_text)
    if normalized_time in {"휴방", "휴뱅"}:
        normalized_time = "휴뱅"

    time_fill = "#f3f4f6" if normalized_time == "휴뱅" else ("#e6f0ff" if normalized_time else "#ffffff")
    requests_payload.append(write_cell_request(sheet_id, base_row, base_col + 1, normalized_time, time_fill))

    padded_entries = list(entries[:3]) + [{"text": "", "categoryKey": ""}] * max(0, 3 - len(entries))
    for offset, entry in zip(CONTENT_ROW_OFFSETS, padded_entries):
        entry_text = normalize_text(entry.get("text"))
        category_key = normalize_text(entry.get("categoryKey"))
        fill_hex = legend_lookup.get(category_key, {}).get("color", "#ffffff") if entry_text else "#ffffff"
        requests_payload.append(write_cell_request(sheet_id, base_row + offset, base_col, entry_text, fill_hex))

    try:
        service.spreadsheets().batchUpdate(
            spreadsheetId=SPREADSHEET_ID,
            body={"requests": requests_payload},
        ).execute()
    except GoogleAPIError as exc:
        raise RuntimeError("Google Sheets 업데이트에 실패했습니다.") from exc


@app.get("/")
def index():
    return send_from_directory(STATIC_DIR, "index.html")


@app.get("/api/bootstrap")
def api_bootstrap():
    try:
        return build_json_response(build_bootstrap_payload())
    except Exception as exc:
        return build_json_response({"error": f"기본 데이터를 불러오지 못했습니다: {exc}"}, 500)


@app.get("/api/calendar/<int:year>/<int:month>")
def api_calendar_month(year: int, month: int):
    try:
        return build_json_response(get_month_or_404(year, month))
    except ValueError:
        return build_json_response({"error": "잘못된 연도 또는 월입니다."}, 400)
    except Exception as exc:
        return build_json_response({"error": f"월간 데이터를 불러오지 못했습니다: {exc}"}, 500)


@app.get("/api/live-status")
def api_live_status():
    return build_json_response(get_live_status_payload())


@app.get("/api/youtube/latest")
def api_youtube_latest():
    return build_json_response(get_youtube_latest_payload())


@app.get("/api/cafe/notice")
def api_cafe_notice():
    try:
        return build_json_response({"item": fetch_cafe_notice_item()})
    except Exception as exc:
        with cafe_notice_lock:
            cached = cafe_notice_cache.get("data")
        if cached:
            return build_json_response({"item": cached, "stale": True})
        return build_json_response({"error": f"공지사항을 불러오지 못했습니다: {exc}"}, 502)


@app.get("/api/songbook")
def api_songbook():
    try:
        snapshot = load_snapshot()
        songbook = build_songbook_payload(snapshot["songbook"])
        items = list(songbook["items"])
        query = normalize_text(request.args.get("q")).lower()
        category = normalize_text(request.args.get("category"))
        artist = normalize_text(request.args.get("artist"))
        if category:
            items = [item for item in items if item["category"] == category]
        if artist:
            items = [item for item in items if item["artist"] == artist]
        if query:
            items = [
                item
                for item in items
                if query in item["title"].lower()
                or query in item["originalTitle"].lower()
                or query in item["artist"].lower()
                or query in item["category"].lower()
            ]
        highlight = random.choice(items) if items else None
        return build_json_response(
            {
                "notice": songbook["notice"],
                "publicLink": songbook["publicLink"],
                "highlight": highlight,
                "items": items,
                "total": len(items),
                "categories": songbook.get("categories", []),
                "artists": songbook.get("artists", []),
                "filters": {"category": category, "artist": artist, "query": query},
            }
        )
    except Exception as exc:
        return build_json_response({"error": f"노래책을 불러오지 못했습니다: {exc}"}, 500)


@app.post("/api/songbook")
def api_songbook_save():
    allowed, message = editor_required()
    if not allowed:
        return build_json_response({"error": message}, 403)

    payload = request.get_json(silent=True) or {}
    snapshot = load_snapshot()
    existing_items = build_songbook_payload(snapshot["songbook"]).get("items") or []
    raw_items = payload.get("items")
    if not isinstance(raw_items, list):
        return build_json_response({"error": "저장할 노래 목록이 올바르지 않습니다."}, 400)

    numeric_numbers = [parse_song_number(item.get("number")) for item in existing_items]
    next_number = max((value for value in numeric_numbers if value is not None), default=0) + 1

    prepared: list[dict[str, Any]] = []
    for index, raw in enumerate(raw_items[:1000], start=1):
        fallback_id = normalize_text((raw or {}).get("id")) or f"song-{index}"
        fallback_number = normalize_text((raw or {}).get("number")) or str(next_number)
        item = sanitize_songbook_item(raw, fallback_number=fallback_number, fallback_id=fallback_id)
        if item:
            if not normalize_text((raw or {}).get("number")):
                next_number += 1
            prepared.append(item)

    if not prepared:
        return build_json_response({"error": "최소 한 곡은 남아 있어야 합니다."}, 400)

    save_songbook_override_items(prepared)
    songbook = build_songbook_payload(snapshot["songbook"])
    return build_json_response(
        {
            "items": songbook["items"],
            "categories": songbook["categories"],
            "artists": songbook["artists"],
            "total": len(songbook["items"]),
        }
    )


@app.get("/api/postit")
def api_postit():
    try:
        return build_json_response(load_postit_payload())
    except Exception as exc:
        return build_json_response({"error": f"포스트잇을 불러오지 못했습니다: {exc}"}, 500)


@app.post("/api/postit")
def api_postit_save():
    allowed, message = postit_editor_required()
    if not allowed:
        return build_json_response({"error": message}, 403)
    payload = request.get_json(silent=True) or {}
    try:
        return build_json_response(save_postit_payload(payload.get("items") or []))
    except Exception as exc:
        return build_json_response({"error": f"포스트잇 저장에 실패했습니다: {exc}"}, 500)


@app.post("/api/auth/google")
def api_auth_google():
    if not google_login_enabled():
        return build_json_response({"error": "Google 로그인 설정이 아직 완료되지 않았습니다."}, 400)
    payload = request.get_json(silent=True) or {}
    credential = normalize_text(payload.get("credential"))
    if not credential:
        return build_json_response({"error": "Google credential 이 비어 있습니다."}, 400)

    try:
        ticket = id_token.verify_oauth2_token(credential, google_requests.Request(), GOOGLE_CLIENT_ID)
    except Exception as exc:
        return build_json_response({"error": f"Google 로그인 검증에 실패했습니다: {exc}"}, 401)

    email = normalize_text(ticket.get("email")).lower()
    if not email:
        return build_json_response({"error": "이메일 정보를 확인하지 못했습니다."}, 400)
    if not ticket.get("email_verified"):
        return build_json_response({"error": "검증된 Google 이메일만 사용할 수 있습니다."}, 403)

    session_payload = {
        "email": email,
        "name": normalize_text(ticket.get("name")) or email.split("@")[0],
        "expiresAt": (now_kst() + timedelta(days=7)).isoformat(),
    }
    cookie_value = build_session_cookie(session_payload)
    response = jsonify(
        {
            "signedIn": True,
            "email": email,
            "displayName": session_payload["name"],
            "canEdit": email in ALLOWED_EDITOR_EMAILS,
            "canPostitEdit": email in ALLOWED_EDITOR_EMAILS,
            "canAdmin": email in ADMIN_EMAILS,
        }
    )
    response.set_cookie(
        SESSION_COOKIE_NAME,
        cookie_value,
        httponly=True,
        samesite="Lax",
        secure=False,
        max_age=7 * 24 * 60 * 60,
    )
    return response


@app.post("/api/auth/logout")
def api_auth_logout():
    response = jsonify({"signedIn": False})
    response.delete_cookie(SESSION_COOKIE_NAME)
    return response


@app.post("/api/admin/day")
def api_admin_day():
    allowed, message = editor_required()
    if not allowed:
        return build_json_response({"error": message}, 403)

    payload = request.get_json(silent=True) or {}
    year = int(payload.get("year") or 0)
    month = int(payload.get("month") or 0)
    day = int(payload.get("day") or 0)
    time_text = normalize_text(payload.get("timeText"))
    day_note_text = normalize_text(payload.get("dayNoteText"))[:1200]
    entries = payload.get("entries") or []

    try:
        month_data = get_month_or_404(year, month)
    except ValueError:
        return build_json_response({"error": "수정할 월 정보를 확인하지 못했습니다."}, 400)

    target_day = next((item for item in month_data["days"] if item["day"] == day), None)
    if not target_day:
        return build_json_response({"error": "해당 날짜를 찾지 못했습니다."}, 404)

    normalized_entries = sanitize_override_entries(entries)

    try:
        normalized_time = normalize_text(format_schedule_badge(time_text))
        if normalized_time in {"휴방", "휴뱅"}:
            normalized_time = "휴뱅"

        update_calendar_day_override(
            year,
            month,
            day,
            normalized_time,
            normalized_entries,
            day_note_text,
            time_source="manual",
        )

        warning = ""
        extra_warning_parts = []
        if writer_enabled() and month_data.get("sheetName") and int(target_day.get("baseRow") or 0) > 0:
            try:
                update_day_on_sheet(
                    month_data["sheetName"],
                    int(target_day["baseRow"]),
                    int(target_day["baseColumn"]),
                    normalized_time,
                    normalized_entries,
                )
                load_snapshot(force=True)
            except Exception as sheet_exc:
                warning = f"사이트에는 저장되었지만 기존 구글 시트 동기화는 실패했습니다: {sheet_exc}"
        if len(normalized_entries) > 3:
            extra_warning_parts.append("4번째 이후 일정은 사이트에만 저장됩니다.")
        if any(normalize_text(item.get("noteText")) for item in normalized_entries):
            extra_warning_parts.append("일정 기타 메모는 사이트에만 저장됩니다.")
        if day_note_text:
            extra_warning_parts.append("날짜 메모는 사이트에만 저장됩니다.")
        if extra_warning_parts:
            warning = " ".join(part for part in [warning, " ".join(extra_warning_parts)] if part).strip()

        updated = get_month_payload(year, month)
        fresh_day = next((item for item in updated["days"] if item["day"] == day), None)
        return build_json_response(
            {
                "ok": True,
                "day": fresh_day,
                "warning": warning,
                "sheetSynced": bool(writer_enabled() and month_data.get("sheetName") and not warning),
            }
        )
    except Exception as exc:
        return build_json_response({"error": f"일정 저장에 실패했습니다: {exc}"}, 500)


@app.get("/static/<path:filename>")
def static_files(filename: str):
    return send_from_directory(STATIC_DIR, filename)


@app.get("/videos/<path:filename>")
def video_files(filename: str):
    return send_from_directory(VIDEOS_DIR, filename)


if __name__ == "__main__":
    app.run(host=HOST, port=PORT, debug=True)
